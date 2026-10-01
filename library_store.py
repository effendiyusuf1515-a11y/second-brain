"""Persistent, local-only document library backed by SQLite and Ollama."""

from __future__ import annotations

import hashlib
import json
import math
import re
import sqlite3
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Generator

import requests

from bibliometrics import COUNTRY_LOCATIONS, infer_document_metadata
from research_core import ExtractedDocument, SourceSection, chunk_document, search_chunks


def _embed_texts(texts: list[str], base_url: str, model: str) -> list[list[float]]:
    vectors: list[list[float]] = []
    for start in range(0, len(texts), 24):
        response = requests.post(
            f"{base_url.rstrip('/')}/api/embed",
            json={"model": model, "input": texts[start : start + 24]},
            timeout=(5, 180),
        )
        response.raise_for_status()
        vectors.extend(response.json()["embeddings"])
    return vectors


def _cosine_similarity(left: list[float], right: list[float]) -> float:
    if len(left) != len(right):
        return -1.0
    numerator = sum(a * b for a, b in zip(left, right))
    left_norm = math.sqrt(sum(value * value for value in left))
    right_norm = math.sqrt(sum(value * value for value in right))
    if not left_norm or not right_norm:
        return 0.0
    return numerator / (left_norm * right_norm)


class ResearchLibrary:
    def __init__(self, data_dir: str | Path) -> None:
        self.data_dir = Path(data_dir)
        self.sources_dir = self.data_dir / "sources"
        self.sources_dir.mkdir(parents=True, exist_ok=True)
        self.database_path = self.data_dir / "library.sqlite3"
        self._initialize()

    @contextmanager
    def _connect(self) -> Generator[sqlite3.Connection, None, None]:
        connection = sqlite3.connect(self.database_path, timeout=30)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        try:
            yield connection
        except BaseException:
            connection.rollback()
            raise
        else:
            connection.commit()
        finally:
            connection.close()

    def _initialize(self) -> None:
        with self._connect() as connection:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS documents (
                    document_id TEXT PRIMARY KEY,
                    filename TEXT NOT NULL,
                    file_type TEXT NOT NULL,
                    source_path TEXT NOT NULL,
                    added_at TEXT NOT NULL,
                    title TEXT NOT NULL DEFAULT '',
                    year INTEGER,
                    country TEXT NOT NULL DEFAULT '',
                    country_source TEXT NOT NULL DEFAULT 'unknown',
                    keywords TEXT NOT NULL DEFAULT '',
                    metadata_version TEXT NOT NULL DEFAULT ''
                );
                CREATE TABLE IF NOT EXISTS chunks (
                    document_id TEXT NOT NULL REFERENCES documents(document_id) ON DELETE CASCADE,
                    ordinal INTEGER NOT NULL,
                    locator TEXT NOT NULL,
                    kind TEXT NOT NULL,
                    text TEXT NOT NULL,
                    embedding_model TEXT NOT NULL DEFAULT '',
                    embedding TEXT NOT NULL DEFAULT '',
                    PRIMARY KEY (document_id, ordinal)
                );
                CREATE INDEX IF NOT EXISTS chunks_embedding_model
                    ON chunks(embedding_model);
                CREATE TABLE IF NOT EXISTS drafts (
                    draft_id TEXT PRIMARY KEY,
                    title TEXT NOT NULL DEFAULT '',
                    research_question TEXT NOT NULL DEFAULT '',
                    research_notes TEXT NOT NULL DEFAULT '',
                    sections_json TEXT NOT NULL DEFAULT '{}',
                    sources_json TEXT NOT NULL DEFAULT '{}',
                    updated_at TEXT NOT NULL
                );
                """
            )
            existing_columns = {
                row["name"] for row in connection.execute("PRAGMA table_info(documents)")
            }
            migrations = {
                "title": "TEXT NOT NULL DEFAULT ''",
                "year": "INTEGER",
                "country": "TEXT NOT NULL DEFAULT ''",
                "country_source": "TEXT NOT NULL DEFAULT 'unknown'",
                "keywords": "TEXT NOT NULL DEFAULT ''",
                "metadata_version": "TEXT NOT NULL DEFAULT ''",
            }
            for column, definition in migrations.items():
                if column not in existing_columns:
                    connection.execute(
                        f"ALTER TABLE documents ADD COLUMN {column} {definition}"
                    )

    def add_document(
        self,
        document: ExtractedDocument,
        original_data: bytes,
        base_url: str,
        embedding_model: str = "",
    ) -> tuple[str, bool, str]:
        document_id = hashlib.sha256(original_data).hexdigest()
        chunks = chunk_document(document)
        metadata = infer_document_metadata(document)
        if not chunks:
            raise ValueError(f"No searchable text found in {document.filename}.")

        safe_name = re.sub(r"[^\w.-]+", "_", Path(document.filename).name).strip("._")
        source_path = self.sources_dir / f"{document_id[:16]}_{safe_name}"
        if not source_path.exists():
            source_path.write_bytes(original_data)

        try:
            with self._connect() as connection:
                existing = connection.execute(
                    "SELECT 1 FROM documents WHERE document_id = ?", (document_id,)
                ).fetchone()
                if existing:
                    return document_id, False, "already indexed"

                connection.execute(
                    """
                    INSERT INTO documents (
                        document_id, filename, file_type, source_path, added_at,
                        title, year, country, country_source, keywords, metadata_version
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        document_id,
                        document.filename,
                        document.file_type,
                        str(source_path),
                        datetime.now(timezone.utc).isoformat(),
                        metadata["title"],
                        metadata["year"],
                        metadata["country"],
                        metadata["country_source"],
                        metadata["keywords"],
                        "1",
                    ),
                )

                vectors: list[list[float]] = []
                vector_model = ""
                embedding_note = "keyword search"
                if embedding_model:
                    try:
                        vectors = _embed_texts(
                            [chunk["text"] for chunk in chunks], base_url, embedding_model
                        )
                        if len(vectors) != len(chunks):
                            raise ValueError("Ollama returned an incomplete embedding batch.")
                        vector_model = embedding_model
                        embedding_note = f"semantic search ({embedding_model})"
                    except (requests.RequestException, KeyError, ValueError) as error:
                        embedding_note = f"keyword search; embedding unavailable: {error}"

                connection.executemany(
                    """
                    INSERT INTO chunks
                        (document_id, ordinal, locator, kind, text, embedding_model, embedding)
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                    """,
                    [
                        (
                            document_id,
                            ordinal,
                            chunk["locator"],
                            chunk["kind"],
                            chunk["text"],
                            vector_model,
                            json.dumps(vectors[ordinal]) if vectors else "",
                        )
                        for ordinal, chunk in enumerate(chunks)
                    ],
                )
        except Exception:
            source_path.unlink(missing_ok=True)
            raise
        return document_id, True, embedding_note

    def contains_document(self, original_data: bytes) -> bool:
        document_id = hashlib.sha256(original_data).hexdigest()
        with self._connect() as connection:
            row = connection.execute(
                "SELECT 1 FROM documents WHERE document_id = ?", (document_id,)
            ).fetchone()
        return row is not None

    def update_metadata(
        self,
        document_id: str,
        *,
        title: str,
        year: int | None,
        country: str,
        keywords: str,
    ) -> None:
        normalized_country = country.strip()
        if normalized_country and normalized_country not in COUNTRY_LOCATIONS:
            raise ValueError("Choose a country with a supported map location.")
        if year is not None and not 1800 <= year <= 2100:
            raise ValueError("Publication year must be between 1800 and 2100.")
        with self._connect() as connection:
            connection.execute(
                """
                UPDATE documents
                SET title = ?, year = ?, country = ?, country_source = ?,
                    keywords = ?, metadata_version = '1'
                WHERE document_id = ?
                """,
                (
                    title.strip(),
                    year,
                    normalized_country,
                    "manually verified" if normalized_country else "unknown",
                    keywords.strip(),
                    document_id,
                ),
            )

    def backfill_metadata(self) -> tuple[int, int]:
        with self._connect() as connection:
            rows = connection.execute(
                """
                SELECT document_id, filename, file_type, source_path,
                       title, year, country, country_source, keywords
                FROM documents
                WHERE metadata_version = ''
                ORDER BY filename
                """
            ).fetchall()

        updated = 0
        failed = 0
        for row in rows:
            try:
                text = ""
                source_path = Path(row["source_path"])
                if row["file_type"] == "pdf":
                    import pymupdf

                    pdf = pymupdf.open(source_path)
                    if pdf.page_count:
                        text = pdf[0].get_text()
                    pdf.close()
                elif row["file_type"] == "docx":
                    from docx import Document

                    docx = Document(source_path)
                    text = "\n".join(
                        paragraph.text for paragraph in docx.paragraphs[:80]
                    )
                elif row["file_type"] == "pptx":
                    from pptx import Presentation

                    presentation = Presentation(source_path)
                    if presentation.slides:
                        text = "\n".join(
                            shape.text
                            for shape in presentation.slides[0].shapes
                            if shape.has_text_frame
                        )

                metadata = infer_document_metadata(
                    ExtractedDocument(
                        row["filename"],
                        row["file_type"],
                        (SourceSection("Page 1", text),),
                    )
                )
                with self._connect() as connection:
                    connection.execute(
                        """
                        UPDATE documents
                        SET title = ?, year = ?, country = ?, country_source = ?,
                            keywords = ?, metadata_version = '1'
                        WHERE document_id = ?
                        """,
                        (
                            row["title"] or metadata["title"],
                            row["year"] or metadata["year"],
                            row["country"] or metadata["country"],
                            row["country_source"]
                            if row["country_source"] != "unknown"
                            else metadata["country_source"],
                            row["keywords"] or metadata["keywords"],
                            row["document_id"],
                        ),
                    )
                updated += 1
            except Exception:
                failed += 1
        return updated, failed

    def documents(self) -> list[dict[str, Any]]:
        with self._connect() as connection:
            rows = connection.execute(
                """
                SELECT d.document_id, d.filename, d.file_type, d.added_at,
                      d.title, d.year, d.country, d.country_source, d.keywords,
                      COUNT(c.ordinal) AS chunk_count
                FROM documents d LEFT JOIN chunks c USING (document_id)
                GROUP BY d.document_id
                ORDER BY d.added_at DESC
                """
            ).fetchall()
        return [dict(row) for row in rows]

    def delete_document(self, document_id: str) -> None:
        with self._connect() as connection:
            row = connection.execute(
                "SELECT source_path FROM documents WHERE document_id = ?", (document_id,)
            ).fetchone()
            if not row:
                return
            connection.execute("DELETE FROM documents WHERE document_id = ?", (document_id,))
        Path(row["source_path"]).unlink(missing_ok=True)

    def search(
        self,
        query: str,
        base_url: str,
        embedding_model: str = "",
        limit: int = 6,
    ) -> tuple[list[dict[str, str]], str]:
        with self._connect() as connection:
            rows = connection.execute(
                """
                SELECT d.filename AS source, c.locator, c.kind, c.text,
                       c.embedding_model, c.embedding
                FROM chunks c JOIN documents d USING (document_id)
                """
            ).fetchall()

        chunks = [
            {key: row[key] for key in ("source", "locator", "kind", "text")}
            for row in rows
        ]
        if embedding_model:
            vector_rows = [row for row in rows if row["embedding_model"] == embedding_model]
            try:
                query_vector = _embed_texts([query], base_url, embedding_model)[0]
                if vector_rows:
                    ranked = sorted(
                        [
                            (
                                _cosine_similarity(
                                    query_vector, json.loads(row["embedding"])
                                ),
                                {
                                    key: row[key]
                                    for key in ("source", "locator", "kind", "text")
                                },
                            )
                            for row in vector_rows
                        ],
                        key=lambda item: item[0],
                        reverse=True,
                    )
                    return [chunk for _, chunk in ranked[:limit]], "semantic"
            except (requests.RequestException, KeyError, IndexError, ValueError):
                pass
        return search_chunks(chunks, query, limit), "keyword"

    def all_chunks(self) -> list[dict[str, str]]:
        with self._connect() as connection:
            rows = connection.execute(
                """
                SELECT d.filename AS source, c.locator, c.kind, c.text
                FROM chunks c JOIN documents d USING (document_id)
                ORDER BY d.filename, c.ordinal
                """
            ).fetchall()
        return [dict(row) for row in rows]

    def load_draft(self, draft_id: str = "main") -> dict[str, Any]:
        with self._connect() as connection:
            row = connection.execute(
                "SELECT * FROM drafts WHERE draft_id = ?", (draft_id,)
            ).fetchone()
        if not row:
            return {
                "title": "",
                "research_question": "",
                "research_notes": "",
                "sections": {},
                "sources": {},
            }
        return {
            "title": row["title"],
            "research_question": row["research_question"],
            "research_notes": row["research_notes"],
            "sections": json.loads(row["sections_json"]),
            "sources": json.loads(row["sources_json"]),
        }

    def save_draft(
        self,
        *,
        title: str,
        research_question: str,
        research_notes: str,
        sections: dict[str, str],
        sources: dict[str, list[str]],
        draft_id: str = "main",
    ) -> None:
        with self._connect() as connection:
            connection.execute(
                """
                INSERT INTO drafts (
                    draft_id, title, research_question, research_notes,
                    sections_json, sources_json, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(draft_id) DO UPDATE SET
                    title = excluded.title,
                    research_question = excluded.research_question,
                    research_notes = excluded.research_notes,
                    sections_json = excluded.sections_json,
                    sources_json = excluded.sources_json,
                    updated_at = excluded.updated_at
                """,
                (
                    draft_id,
                    title,
                    research_question,
                    research_notes,
                    json.dumps(sections, ensure_ascii=False),
                    json.dumps(sources, ensure_ascii=False),
                    datetime.now(timezone.utc).isoformat(),
                ),
            )
