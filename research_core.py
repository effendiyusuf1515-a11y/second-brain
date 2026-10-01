"""Local document extraction and source-aware retrieval helpers."""

from __future__ import annotations

import base64
import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path

import requests
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".xlsx",
    ".pptx",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".tif",
    ".tiff",
    ".txt",
    ".csv",
}
MAX_VISION_IMAGES = 2


@dataclass(frozen=True)
class SourceSection:
    locator: str
    text: str
    kind: str = "text"


@dataclass(frozen=True)
class ExtractedDocument:
    filename: str
    file_type: str
    sections: tuple[SourceSection, ...]


def _vision_description(
    image_data: bytes, ollama_url: str, vision_model: str, locator: str
) -> str:
    response = requests.post(
        f"{ollama_url.rstrip('/')}/api/chat",
        json={
            "model": vision_model,
            "stream": False,
            "messages": [
                {
                    "role": "user",
                    "content": (
                        f"Analyze this research document image ({locator}). Transcribe "
                        "visible text and explain tables, charts, diagrams, axes, units, "
                        "and relationships. Separate observations from uncertainty; "
                        "do not guess missing values."
                    ),
                    "images": [base64.b64encode(image_data).decode("ascii")],
                }
            ],
            "options": {"temperature": 0.1},
        },
        timeout=(5, 240),
    )
    response.raise_for_status()
    return response.json()["message"]["content"].strip()


def _clean_pdf_text(text: str) -> str:
    return text.encode("utf-8", errors="replace").decode("utf-8").strip()


def _image_section(
    image_data: bytes,
    locator: str,
    ollama_url: str,
    vision_model: str,
) -> SourceSection:
    description = _vision_description(image_data, ollama_url, vision_model, locator)
    return SourceSection(locator, description, "visual")


def _extract_pdf(
    filename: str,
    data: bytes,
    ollama_url: str,
    vision_model: str,
    inspect_visuals: bool,
) -> list[SourceSection]:
    reader = PdfReader(io.BytesIO(data))
    pages = [
        (number, _clean_pdf_text(page.extract_text() or ""))
        for number, page in enumerate(reader.pages, 1)
    ]
    visual_pages = {number for number, text in pages if len(text) < 80}
    if inspect_visuals:
        visual_pages.update(number for number, _ in pages)

    rendered_pages: dict[int, bytes] = {}
    if visual_pages and vision_model:
        try:
            import pymupdf
        except ImportError as error:
            raise RuntimeError(
                "PDF visual analysis requires PyMuPDF. Install the project requirements."
            ) from error
        pdf = pymupdf.open(stream=data, filetype="pdf")
        for page_number in sorted(visual_pages)[:MAX_VISION_IMAGES]:
            pixmap = pdf[page_number - 1].get_pixmap(
            matrix=pymupdf.Matrix(1.5, 1.5), alpha=False
            )
            rendered_pages[page_number] = pixmap.tobytes("jpeg")
        pdf.close()

    sections = []
    for page_number, text in pages:
        parts = [text] if text else []
        if page_number in rendered_pages:
            visual_text = _vision_description(
                rendered_pages[page_number], ollama_url, vision_model, f"page {page_number}"
            )
            if visual_text:
                parts.append(f"Visual analysis: {visual_text}")
        if parts:
            sections.append(SourceSection(f"Page {page_number}", "\n\n".join(parts)))
    if not sections:
        raise ValueError(f"No text or visual content could be extracted from {filename}.")
    return sections


def _extract_docx(
    data: bytes, ollama_url: str, vision_model: str
) -> list[SourceSection]:
    from docx import Document

    document = Document(io.BytesIO(data))
    sections: list[SourceSection] = []
    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]
    if paragraphs:
        sections.append(SourceSection("Document text", "\n\n".join(paragraphs)))
    for table_number, table in enumerate(document.tables, 1):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        sections.append(SourceSection(f"Table {table_number}", "\n".join(rows), "table"))
    if vision_model:
        images = [
            part.blob
            for part in document.part.related_parts.values()
            if part.content_type.startswith("image/")
        ]
        for image_number, image_data in enumerate(images[:MAX_VISION_IMAGES], 1):
            sections.append(
                _image_section(
                    image_data,
                    f"embedded image {image_number}",
                    ollama_url,
                    vision_model,
                )
            )
    return sections


def _chart_reference_values(reference: object, worksheet: object) -> list[str]:
    formula = getattr(reference, "f", "")
    if not formula:
        return []
    target_sheet = worksheet
    cell_range = formula
    if "!" in formula:
        sheet_name, cell_range = formula.rsplit("!", 1)
        sheet_name = sheet_name.strip("'").replace("''", "'")
        try:
            target_sheet = worksheet.parent[sheet_name]
        except (AttributeError, KeyError, TypeError):
            return []
    try:
        from openpyxl.utils.cell import range_boundaries

        min_column, min_row, max_column, max_row = range_boundaries(
            cell_range.replace("$", "")
        )
        values = []
        for row in target_sheet.iter_rows(
            min_row=min_row,
            max_row=max_row,
            min_col=min_column,
            max_col=max_column,
            values_only=True,
        ):
            values.extend(str(value) for value in row if value is not None)
        return values
    except (TypeError, ValueError):
        return []


def _chart_summary(chart: object, worksheet: object | None = None) -> str:
    summaries = [f"Chart type: {type(chart).__name__}"]
    for series_number, series in enumerate(getattr(chart, "series", []), 1):
        values = list(getattr(series, "values", []) or [])
        categories = list(getattr(series, "category_values", []) or [])
        if worksheet is not None:
            value_reference = getattr(getattr(series, "val", None), "numRef", None)
            category = getattr(series, "cat", None)
            category_reference = getattr(category, "strRef", None) or getattr(
                category, "numRef", None
            )
            if not values:
                values = _chart_reference_values(value_reference, worksheet)
            if not categories:
                categories = _chart_reference_values(category_reference, worksheet)
        name = getattr(series, "name", None) or f"Series {series_number}"
        summaries.append(
            f"{name}: "
            f"categories={categories[:80] if categories else 'not available'}; "
            f"values={values[:80] if values else 'not available'}"
        )
    return "\n".join(summaries)


def _extract_xlsx(
    data: bytes, ollama_url: str, vision_model: str
) -> list[SourceSection]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    sections: list[SourceSection] = []
    for worksheet in workbook.worksheets:
        rows = []
        for row in worksheet.iter_rows(values_only=True):
            values = [
                str(value).strip()
                for value in row
                if value is not None and str(value).strip()
            ]
            if values:
                rows.append(" | ".join(values))
        if rows:
            sections.append(
                SourceSection(f"Sheet: {worksheet.title}", "\n".join(rows), "table")
            )
        for chart_number, chart in enumerate(worksheet._charts, 1):
            sections.append(
                SourceSection(
                    f"Sheet: {worksheet.title}, chart {chart_number}",
                    _chart_summary(chart, worksheet),
                    "chart",
                )
            )
        if vision_model:
            for image_number, image in enumerate(
                worksheet._images[:MAX_VISION_IMAGES], 1
            ):
                sections.append(
                    _image_section(
                        image._data(),
                        f"sheet {worksheet.title}, embedded image {image_number}",
                        ollama_url,
                        vision_model,
                    )
                )
    workbook.close()
    return sections


def _extract_pptx(
    data: bytes, ollama_url: str, vision_model: str
) -> list[SourceSection]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(io.BytesIO(data))
    sections: list[SourceSection] = []
    for slide_number, slide in enumerate(presentation.slides, 1):
        text_parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                text_parts.append(shape.text.strip())
            if shape.has_table:
                for row in shape.table.rows:
                    text_parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            if shape.has_chart:
                text_parts.append(_chart_summary(shape.chart))
            if vision_model and shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                sections.append(
                    _image_section(
                        shape.image.blob,
                        f"slide {slide_number}, image",
                        ollama_url,
                        vision_model,
                    )
                )
        if text_parts:
            sections.append(SourceSection(f"Slide {slide_number}", "\n".join(text_parts)))
    return sections


def extract_document(
    filename: str,
    data: bytes,
    ollama_url: str = "http://localhost:11434",
    vision_model: str = "",
    inspect_visuals: bool = False,
) -> ExtractedDocument:
    extension = Path(filename).suffix.casefold()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {extension or 'no extension'}")

    if extension == ".pdf":
        sections = _extract_pdf(filename, data, ollama_url, vision_model, inspect_visuals)
    elif extension == ".docx":
        sections = _extract_docx(data, ollama_url, vision_model)
    elif extension == ".xlsx":
        sections = _extract_xlsx(data, ollama_url, vision_model)
    elif extension == ".pptx":
        sections = _extract_pptx(data, ollama_url, vision_model)
    elif extension in {".txt", ".csv"}:
        text = data.decode("utf-8-sig", errors="replace")
        if extension == ".csv":
            rows = csv.reader(io.StringIO(text))
            text = "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
        sections = [SourceSection("Document text", text.strip())] if text.strip() else []
    else:
        if not vision_model:
            raise ValueError("Choose a Qwen-VL model to analyze image files.")
        sections = [_image_section(data, "Image", ollama_url, vision_model)]

    if not sections:
        raise ValueError(f"No readable content found in {filename}.")
    return ExtractedDocument(filename, extension.lstrip("."), tuple(sections))


def chunk_document(
    document: ExtractedDocument, chunk_size: int = 1600, overlap: int = 180
) -> list[dict[str, str]]:
    chunks: list[dict[str, str]] = []
    for section in document.sections:
        text = re.sub(r"\s+", " ", section.text).strip()
        start = 0
        while start < len(text):
            end = min(start + chunk_size, len(text))
            if end < len(text):
                boundary = text.rfind(" ", start + chunk_size // 2, end)
                if boundary > start:
                    end = boundary
            content = text[start:end].strip()
            if content:
                chunks.append(
                    {
                        "source": document.filename,
                        "locator": section.locator,
                        "kind": section.kind,
                        "text": content,
                    }
                )
            if end >= len(text):
                break
            start = max(end - overlap, start + 1)
    return chunks


def search_chunks(
    chunks: list[dict[str, str]], query: str, limit: int = 6
) -> list[dict[str, str]]:
    stop_words = {"dan", "yang", "untuk", "dari", "the", "and", "for", "with"}
    terms = {
        term.casefold()
        for term in re.findall(r"[\w-]{3,}", query)
        if term.casefold() not in stop_words
    }
    if not terms:
        return chunks[:limit]

    ranked = []
    for chunk in chunks:
        words = re.findall(r"[\w-]{3,}", chunk["text"].casefold())
        score = sum(min(words.count(term), 4) for term in terms)
        if score:
            ranked.append((score, chunk))
    ranked.sort(key=lambda item: (-item[0], item[1]["source"], item[1]["locator"]))
    return [chunk for _, chunk in ranked[:limit]]
