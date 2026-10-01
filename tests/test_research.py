from __future__ import annotations

import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pymupdf
import requests
from docx import Document
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from PIL import Image
from pptx import Presentation

from bibliometrics import bibliometric_summary, filter_documents, infer_document_metadata
import library_store
from library_store import ResearchLibrary
from research_core import (
    ExtractedDocument,
    SourceSection,
    chunk_document,
    extract_document,
    search_chunks,
    _clean_pdf_text,
)


class ResearchCoreTests(unittest.TestCase):
    def test_infers_title_year_country_and_author_keywords(self) -> None:
        document = ExtractedDocument(
            "study.pdf",
            "pdf",
            (
                SourceSection(
                    "Page 1",
                    "Model assessment of bridges\n"
                    "University of Sannio, Benevento, Italy\n"
                    "Available online 17 November 2022\n"
                    "Abstract\nBridge response.\n"
                    "Keywords: finite element model updating; dynamic load test",
                ),
            ),
        )

        metadata = infer_document_metadata(document)

        self.assertEqual(metadata["title"], "Model assessment of bridges")
        self.assertEqual(metadata["year"], 2022)
        self.assertEqual(metadata["country"], "Italy")
        self.assertIn("finite element model updating", metadata["keywords"])

    def test_prefers_first_country_in_multiple_affiliations(self) -> None:
        document = ExtractedDocument(
            "multinational.pdf",
            "pdf",
            (
                SourceSection(
                    "Page 1",
                    "Bridge study\nFirst author, Nanjing, China; second author, Yokohama, Japan\n"
                    "Abstract\nThe study evaluates bridge response.",
                ),
            ),
        )

        metadata = infer_document_metadata(document)

        self.assertEqual(metadata["country"], "China")

    def test_filters_by_full_text_country_and_year(self) -> None:
        documents = [
            {
                "filename": "italy.pdf",
                "title": "Bridge assessment",
                "keywords": "finite element; static test",
                "country": "Italy",
                "year": 2022,
            },
            {
                "filename": "unknown.pdf",
                "title": "Bridge notes",
                "keywords": "dynamic test",
                "country": "",
                "year": None,
            },
        ]
        chunks = [
            {"source": "italy.pdf", "locator": "Page 4", "kind": "text", "text": "Kriging finite element updating"},
            {"source": "unknown.pdf", "locator": "Page 1", "kind": "text", "text": "dynamic bridge response"},
        ]

        filtered = filter_documents(
            documents,
            chunks,
            query="Kriging updating",
            countries=["Italy"],
            year_range=(2020, 2024),
        )

        self.assertEqual([document["filename"] for document in filtered], ["italy.pdf"])

    def test_bibliometric_summary_labels_inferred_keyword_source(self) -> None:
        summary = bibliometric_summary(
            [{"filename": "bridge.pdf", "title": "Dynamic bridge response", "year": 2022, "country": "Italy"}],
            [],
        )

        self.assertEqual(summary["keyword_source"], "terms inferred from titles")
        self.assertEqual(summary["country_counts"], [("Italy", 1)])

    def test_bibliometric_keywords_ignore_extracted_section_headers(self) -> None:
        summary = bibliometric_summary(
            [{"filename": "bridge.pdf", "title": "Bridge study", "keywords": "a b s t r a c t; finite element"}],
            [],
        )

        self.assertEqual(summary["keyword_counts"], [("finite element", 1)])

    def test_cleans_invalid_surrogate_characters_from_pdf_text(self) -> None:
        self.assertEqual(_clean_pdf_text("bridge \udefe test"), "bridge ? test")

    def test_extracts_word_paragraphs_and_tables(self) -> None:
        source = Document()
        source.add_paragraph("Bridge load testing")
        table = source.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Capacity"
        table.cell(0, 1).text = "145 kN"
        buffer = BytesIO()
        source.save(buffer)

        result = extract_document("bridge.docx", buffer.getvalue())

        self.assertIn("Bridge load testing", result.sections[0].text)
        self.assertEqual(result.sections[1].locator, "Table 1")
        self.assertIn("145 kN", result.sections[1].text)

    def test_extracts_spreadsheet_cells_and_chart_values(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Load tests"
        sheet.append(["Test", "Capacity"])
        sheet.append(["A", 120])
        sheet.append(["B", 145])
        chart = LineChart()
        chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
        sheet.add_chart(chart, "D2")
        buffer = BytesIO()
        workbook.save(buffer)

        result = extract_document("bridge-data.xlsx", buffer.getvalue())
        chart_text = " ".join(section.text for section in result.sections if section.kind == "chart")

        self.assertIn("120", chart_text)
        self.assertIn("145", chart_text)
        self.assertIn("A", chart_text)
        self.assertIn("B", chart_text)

    def test_extracts_presentation_slide_text(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = "Dynamic bridge test"
        buffer = BytesIO()
        presentation.save(buffer)

        result = extract_document("bridge.pptx", buffer.getvalue())

        self.assertIn("Dynamic bridge test", result.sections[0].text)
        self.assertEqual(result.sections[0].locator, "Slide 1")

    def test_routes_scanned_pdf_page_to_vision_model(self) -> None:
        image_buffer = BytesIO()
        Image.new("RGB", (32, 32), "white").save(image_buffer, format="PNG")
        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_image(
            pymupdf.Rect(0, 0, 200, 200), stream=image_buffer.getvalue()
        )
        pdf_bytes = pdf.tobytes()
        pdf.close()

        with patch(
            "research_core._vision_description", return_value="Detected a chart"
        ) as vision_call:
            result = extract_document(
                "scanned.pdf", pdf_bytes, vision_model="qwen3-vl:8b"
            )

        vision_call.assert_called_once()
        self.assertEqual(result.sections[0].locator, "Page 1")
        self.assertIn("Detected a chart", result.sections[0].text)

    def test_limits_automatic_vision_to_two_pdf_pages(self) -> None:
        pdf = pymupdf.open()
        for _ in range(4):
            pdf.new_page()
        pdf_bytes = pdf.tobytes()
        pdf.close()

        with patch(
            "research_core._vision_description", return_value="Visual page"
        ) as vision_call:
            result = extract_document(
                "long-scan.pdf",
                pdf_bytes,
                vision_model="qwen3-vl:8b",
                inspect_visuals=True,
            )

        self.assertEqual(vision_call.call_count, 2)
        self.assertEqual(len(result.sections), 2)

    def test_chunks_and_keyword_results_keep_source_locator(self) -> None:
        document = ExtractedDocument(
            "bridge.pdf",
            "pdf",
            (SourceSection("Page 7", "Bridge load testing results"),),
        )

        chunks = chunk_document(document)
        hits = search_chunks(chunks, "bridge load")

        self.assertEqual(hits[0]["source"], "bridge.pdf")
        self.assertEqual(hits[0]["locator"], "Page 7")


class ResearchLibraryTests(unittest.TestCase):
    def test_saves_and_restores_draft_after_library_reopen(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            first_instance = ResearchLibrary(Path(directory))
            first_instance.save_draft(
                title="Bridge assessment",
                research_question="How do static and dynamic tests compare?",
                research_notes="Use only verified results.",
                sections={"Tinjauan pustaka": "Source-grounded draft."},
                sources={"Tinjauan pustaka": ["paper.pdf · Page 3"]},
            )

            restored = ResearchLibrary(Path(directory)).load_draft()

            self.assertEqual(restored["title"], "Bridge assessment")
            self.assertEqual(restored["research_question"], "How do static and dynamic tests compare?")
            self.assertEqual(restored["research_notes"], "Use only verified results.")
            self.assertEqual(restored["sections"]["Tinjauan pustaka"], "Source-grounded draft.")
            self.assertEqual(restored["sources"]["Tinjauan pustaka"], ["paper.pdf · Page 3"])

    def test_stores_detected_and_manually_verified_bibliographic_metadata(self) -> None:
        document = ExtractedDocument(
            "bridge-study.pdf",
            "pdf",
            (
                SourceSection(
                    "Page 1",
                    "Static and dynamic bridge load tests\n"
                    "University of Sannio, Benevento, Italy\n"
                    "Available online 2023\nAbstract\n"
                    "Keywords: finite element model; bridge load rating",
                ),
            ),
        )
        with tempfile.TemporaryDirectory() as directory:
            library = ResearchLibrary(Path(directory))
            document_id, added, _ = library.add_document(
                document, b"bridge-study-source", "http://localhost:11434"
            )
            detected = library.documents()[0]
            self.assertTrue(added)
            self.assertEqual(detected["country"], "Italy")
            self.assertEqual(detected["country_source"], "detected from first-page affiliation text")
            self.assertEqual(detected["year"], 2023)

            library.update_metadata(
                document_id,
                title="Verified bridge study",
                year=2024,
                country="China",
                keywords="load test; finite element",
            )

            corrected = library.documents()[0]
            self.assertEqual(corrected["title"], "Verified bridge study")
            self.assertEqual(corrected["country"], "China")
            self.assertEqual(corrected["country_source"], "manually verified")
            self.assertEqual(corrected["year"], 2024)
            self.assertEqual(corrected["keywords"], "load test; finite element")

    def test_persists_deduplicates_semantically_searches_and_deletes(self) -> None:
        document = ExtractedDocument(
            "bridge.pdf",
            "pdf",
            (
                SourceSection("Page 7", "Bridge load testing results"),
                SourceSection("Page 8", "Bridge load test comparisons"),
            ),
        )
        with tempfile.TemporaryDirectory() as directory:
            library = ResearchLibrary(Path(directory))
            with patch.object(
                library_store,
                "_embed_texts",
                side_effect=[[[1.0, 0.0], [1.0, 0.0]], [[1.0, 0.0]]],
            ):
                document_id, added, mode = library.add_document(
                    document, b"sample-source", "http://localhost:11434", "test-embed"
                )
                self.assertTrue(library.contains_document(b"sample-source"))
                hits, retrieval = library.search(
                    "bridge load", "http://localhost:11434", "test-embed"
                )
                duplicate_id, duplicate_added, _ = library.add_document(
                    document, b"sample-source", "http://localhost:11434"
                )

            self.assertTrue(added)
            self.assertEqual(mode, "semantic search (test-embed)")
            self.assertEqual(retrieval, "semantic")
            self.assertEqual(len(hits), 2)
            self.assertEqual(duplicate_id, document_id)
            self.assertFalse(duplicate_added)
            self.assertEqual(library.documents()[0]["filename"], "bridge.pdf")

            library.delete_document(document_id)

            self.assertEqual(library.documents(), [])

    def test_missing_embedding_falls_back_to_keyword_search(self) -> None:
        document = ExtractedDocument(
            "bridge.pdf", "pdf", (SourceSection("Page 2", "Bridge load data"),)
        )
        with tempfile.TemporaryDirectory() as directory:
            library = ResearchLibrary(Path(directory))
            with patch.object(
                library_store, "_embed_texts", side_effect=requests.HTTPError("model missing")
            ):
                _, added, note = library.add_document(
                    document, b"source", "http://localhost:11434", "missing-model"
                )
                hits, retrieval = library.search(
                    "bridge load", "http://localhost:11434", "missing-model"
                )

            self.assertTrue(added)
            self.assertIn("keyword search", note)
            self.assertEqual(retrieval, "keyword")
            self.assertEqual(hits[0]["locator"], "Page 2")


if __name__ == "__main__":
    unittest.main()
