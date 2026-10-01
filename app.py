import html
import io
import re
from pathlib import Path

import requests
import streamlit as st
import pandas as pd
import altair as alt

from bibliometrics import COUNTRY_LOCATIONS, bibliometric_summary, filter_documents
from library_store import ResearchLibrary
from research_core import SUPPORTED_EXTENSIONS, extract_document


DEFAULT_MODEL = "freehuntx/qwen3-coder:8b"
MAX_CONTEXT_CHARS = 7500
def ask_ollama(
    base_url: str,
    model: str,
    keep_alive: int | str,
    pdf_name: str,
    context: str,
    history: list[dict[str, str]],
    question: str,
    answer_language: str,
) -> str:
    system_prompt = (
        f"You are a careful research assistant. Answer in {answer_language}. "
        "Use only the supplied excerpts and user-provided research notes. Do not invent "
        "facts, methods, results, sample sizes, or citations. Cite each important claim "
        "with the exact source label in square brackets, including its locator. If the "
        "evidence is insufficient, say what is missing. Treat all generated text as a "
        "draft that needs researcher review.\n\n"
        f"Workspace: {pdf_name}\n\nRetrieved source excerpts:\n{context}"
    )
    messages = [{"role": "system", "content": system_prompt}]
    messages.extend(
        {"role": item["role"], "content": item["content"]}
        for item in history[-4:]
    )
    messages.append({"role": "user", "content": question})

    response = requests.post(
        f"{base_url}/api/chat",
        json={
            "model": model,
            "messages": messages,
            "stream": False,
            "keep_alive": keep_alive,
            "options": {"temperature": 0.2, "num_ctx": 8192},
        },
        timeout=(5, 240),
    )
    response.raise_for_status()
    return response.json()["message"]["content"].strip()


ROOT = Path(__file__).resolve().parent
DATA_DIR = ROOT / "data"
DEFAULT_TEXT_MODEL = "qwen3:8b"
DEFAULT_VISION_MODEL = "qwen3-vl:8b"
DEFAULT_EMBEDDING_MODEL = "nomic-embed-text"
SECTION_ORDER = [
    "Abstract",
    "Introduction",
    "Literature review",
    "Methods",
    "Results",
    "Discussion",
    "Conclusion",
]
LEGACY_SECTION_NAMES = {
    "Abstrak": "Abstract",
    "Pendahuluan": "Introduction",
    "Tinjauan pustaka": "Literature review",
    "Metode": "Methods",
    "Hasil": "Results",
    "Pembahasan": "Discussion",
    "Kesimpulan": "Conclusion",
}


@st.cache_resource
def get_library(data_dir: str) -> ResearchLibrary:
    return ResearchLibrary(data_dir)


@st.cache_data(ttl=15, show_spinner=False)
def get_ollama_models(base_url: str) -> tuple[str, ...]:
    try:
        response = requests.get(f"{base_url.rstrip('/')}/api/tags", timeout=3)
        response.raise_for_status()
        return tuple(model["name"] for model in response.json().get("models", []))
    except (requests.RequestException, KeyError, TypeError, ValueError):
        return ()


def format_context(chunks: list[dict[str, str]]) -> str:
    parts = []
    used_chars = 0
    for chunk in chunks:
        label = f"[{chunk['source']} | {chunk['locator']}]"
        part = f"{label}\n{chunk['text']}"
        available = MAX_CONTEXT_CHARS - used_chars
        if available <= 0:
            break
        parts.append(part[:available])
        used_chars += len(part)
    return "\n\n".join(parts)


def make_docx(title: str, sections: dict[str, str], sources: dict[str, list[str]]) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(title.strip() or "Research paper draft", 0)
    document.add_paragraph("Working draft. Verify all claims and references before publication.")
    for section_name in SECTION_ORDER:
        text = sections.get(section_name, "").strip()
        if not text:
            continue
        document.add_heading(section_name, level=1)
        for paragraph in re.split(r"\n\s*\n", text):
            if paragraph.strip():
                document.add_paragraph(paragraph.strip())
        citations = sources.get(section_name, [])
        if citations:
            paragraph = document.add_paragraph()
            paragraph.add_run("Sources to verify: ").bold = True
            paragraph.add_run("; ".join(citations))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


st.set_page_config(
    page_title="Second Brain | Research workspace",
    page_icon=":material/psychology:",
    layout="wide",
)
THEME_PALETTES = {
        "light": {
                "canvas": "#F7F5EF", "surface": "#FFFEFA", "surface_alt": "#EFEEE7",
                "ink": "#263238", "muted": "#69767A", "line": "#DEDCD2",
                "sidebar": "#E9ECE5", "sidebar_ink": "#243733", "accent": "#BB5A3C",
                "accent_soft": "#F4E4DB", "teal": "#2E766F", "field": "#FFFEFA",
                "paper_ink": "#344247", "button_ink": "#FFFEFA",
        },
        "dark": {
                "canvas": "#171E22", "surface": "#20292D", "surface_alt": "#263236",
                "ink": "#EEF0E9", "muted": "#A7B2AE", "line": "#394549",
                "sidebar": "#101719", "sidebar_ink": "#E7ECE6", "accent": "#F08A69",
                "accent_soft": "#3A2925", "teal": "#79C0B5", "field": "#263236",
                "paper_ink": "#E6E8E1", "button_ink": "#171E22",
        },
}
appearance_mode = st.session_state.get("appearance_mode")
active_theme = (
    appearance_mode.casefold()
    if appearance_mode
    else st.context.theme.type if st.context.theme.type in THEME_PALETTES else "light"
)
if "appearance_mode" not in st.session_state:
    st.session_state.appearance_mode = "Dark" if active_theme == "dark" else "Light"
palette = THEME_PALETTES[active_theme]
theme_variables = "".join(f"--{name}:{value};" for name, value in palette.items())
st.markdown(
        "<style>:root{" + theme_variables + "}"
        + """
        html, body, [class*="css"] { font-family: 'Aptos', 'Segoe UI', sans-serif; letter-spacing: 0; }
        [data-testid="stAppViewContainer"] { background:var(--canvas); color:var(--ink); }
        [data-testid="stHeader"] { background:transparent; }
        [data-testid="stSidebar"] { background:var(--sidebar); border-right:1px solid var(--line); }
        [data-testid="stSidebar"] * { color:var(--sidebar-ink); }
        [data-testid="stSidebar"] [data-testid="stCaptionContainer"] * { color:var(--muted); }
        [data-testid="stSidebar"] [data-testid="stWidgetLabel"] p { color:var(--sidebar-ink); }
        [data-testid="stSidebar"] [data-testid="stTextInputRootElement"] { background:var(--field) !important; border:1px solid var(--line) !important; }
        [data-testid="stSidebar"] [data-testid="stTextInputRootElement"] input { background:transparent !important; color:var(--ink) !important; }
        [data-testid="stSidebar"] [data-baseweb="select"] > div { background:var(--field) !important; color:var(--ink) !important; }
        [data-testid="stSidebar"] [data-testid="stRadio"] label { padding:.42rem .3rem; border-radius:6px; }
        .block-container { max-width:1500px; padding-top:1.2rem; padding-bottom:4rem; }
        h1, h2, h3 { color:var(--ink); letter-spacing:0; }
        h1, h2 { font-family:'Cambria', Georgia, serif; font-weight:600; }
        h1 { font-size:2.5rem; line-height:1.1; }
        h2 { font-size:1.8rem; }
        p, label, div { letter-spacing:0; }
        [data-testid="stMetric"] { background:var(--surface); border:1px solid var(--line); padding:.9rem 1rem; border-radius:6px; }
        [data-testid="stMetricLabel"] p { color:var(--muted); font-size:.78rem; }
        [data-testid="stMetricValue"] { color:var(--ink); font-size:1.65rem; }
        .brand-lockup { display:flex; align-items:center; gap:.65rem; padding:.45rem 0 1.1rem; }
        .brand-mark { display:grid; place-items:center; width:2.2rem; height:2.2rem; border-radius:6px; background:var(--accent); color:var(--button-ink); font-family:Georgia,serif; font-size:1.35rem; font-weight:700; }
        .brand-name { color:var(--sidebar-ink); font-size:1.03rem; font-weight:700; }
        .sidebar-note { color:var(--muted); font-size:.76rem; line-height:1.5; }
        .eyebrow { color:var(--teal); font-size:.71rem; font-weight:700; letter-spacing:.08em; text-transform:uppercase; }
        .hero-copy { max-width:780px; color:var(--muted); font-size:1rem; line-height:1.65; }
        .hero-band { padding:.55rem 0 1.35rem; border-bottom:1px solid var(--line); margin-bottom:1.25rem; }
        .assistant-panel { background:var(--surface); border:1px solid var(--line); border-left:4px solid var(--accent); padding:1.2rem 1.25rem; border-radius:7px; }
        .section-rule { height:1px; background:var(--line); margin:1.2rem 0; }
        .paper-preview { min-height:460px; background:var(--surface); color:var(--paper-ink); border:1px solid var(--line); border-top:5px solid var(--accent); padding:2.2rem 2rem; box-shadow:0 12px 32px rgba(10,20,24,.12); border-radius:6px; }
        .paper-preview h3 { font-family:'Cambria',Georgia,serif; color:var(--ink); font-size:1.5rem; margin:0 0 1.4rem; }
        .paper-preview p { color:var(--paper-ink); font-family:Georgia,serif; font-size:.97rem; line-height:1.85; white-space:pre-wrap; }
        .paper-meta { color:var(--muted); font-size:.68rem; letter-spacing:.06em; text-transform:uppercase; }
        .stButton > button, .stDownloadButton > button { border-radius:5px; font-weight:600; min-height:2.5rem; }
        .stButton > button[kind="primary"] { background:var(--accent); border-color:var(--accent); color:var(--button-ink); }
        [data-testid="stFileUploader"] { background:var(--surface); border:1px dashed var(--muted); border-radius:6px; padding:.6rem; }
        [data-testid="stExpander"] { border-color:var(--line); border-radius:6px; }
        [data-testid="stChatMessage"] { background:var(--surface); border:1px solid var(--line); border-radius:6px; }
        @media (max-width:760px) {
            .block-container { padding:.8rem 1rem 3rem; }
            h1 { font-size:2rem; }
            [data-testid="stMetric"] { padding:.7rem; }
            .paper-preview { padding:1.3rem 1.15rem; }
        }
        </style>
        """,
        unsafe_allow_html=True,
)

library = get_library(str(DATA_DIR))
documents = library.documents()
all_chunks = library.all_chunks()
passage_count = sum(document["chunk_count"] for document in documents)


def select_workspace(target_page: str) -> None:
    st.session_state.workspace_page = target_page


def sync_draft_editor(section_name: str) -> None:
    section_key = f"draft-editor-{section_name}"
    st.session_state.draft_sections[section_name] = st.session_state.get(section_key, "")
    save_current_draft()


def save_current_draft() -> None:
    library.save_draft(
        title=st.session_state.get("draft_title", ""),
        research_question=st.session_state.get("research_question", ""),
        research_notes=st.session_state.get("research_notes", ""),
        sections=st.session_state.get("draft_sections", {}),
        sources=st.session_state.get("draft_sources", {}),
    )


with st.sidebar:
    st.markdown(
        '<div class="brand-lockup"><span class="brand-mark">SB</span>'
        '<span class="brand-name">Second Brain</span></div>',
        unsafe_allow_html=True,
    )
    legacy_pages = {
        "Ruang kerja": "Home",
        "Beranda": "Home",
        "Perpustakaan": "Library",
        "Tanya sumber": "Ask library",
        "Draft jurnal": "Page builder",
    }
    current_page = st.session_state.get("workspace_page")
    if current_page in legacy_pages:
        st.session_state.workspace_page = legacy_pages[current_page]
    page = st.radio(
        "Workspace",
        ["Home", "Library", "Ask library", "Page builder"],
        key="workspace_page",
        label_visibility="collapsed",
    )
    st.markdown('<div class="section-rule"></div>', unsafe_allow_html=True)
    st.segmented_control(
        "Appearance",
        ["Light", "Dark"],
        key="appearance_mode",
        selection_mode="single",
        width="stretch",
    )
    st.markdown("**LOCAL MODELS**")
    ollama_url = st.text_input("Ollama address", "http://localhost:11434").strip().rstrip("/")
    text_model = st.text_input("Writing model", DEFAULT_TEXT_MODEL).strip()
    vision_model = st.text_input("Vision model", DEFAULT_VISION_MODEL).strip()
    embedding_model = st.text_input("Embedding model", DEFAULT_EMBEDDING_MODEL).strip()
    keep_alive_label = st.selectbox("Model memory", ["Unload after each answer", "Keep loaded for 5 minutes"])
    keep_alive: int | str = 0 if keep_alive_label == "Unload after each answer" else "5m"
    if st.button("Test Ollama connection", width="stretch"):
        try:
            response = requests.get(f"{ollama_url}/api/tags", timeout=5)
            response.raise_for_status()
            st.session_state.model_check = "connected"
            st.session_state.available_models = [
                item["name"] for item in response.json().get("models", [])
            ]
        except requests.RequestException as error:
            st.session_state.model_check = str(error)
    available_models = get_ollama_models(ollama_url)
    if st.session_state.get("model_check") == "connected" or available_models:
        st.markdown('<p class="sidebar-note">Ollama is connected. Documents are processed at this endpoint.</p>', unsafe_allow_html=True)
        if embedding_model and embedding_model not in available_models:
            st.markdown(
                '<p class="sidebar-note">Embedding model not detected; keyword search is active until one is installed.</p>',
                unsafe_allow_html=True,
            )
    elif st.session_state.get("model_check"):
        st.error(f"Cannot reach Ollama: {st.session_state.model_check}")
    else:
        st.markdown('<p class="sidebar-note">Start Ollama to use AI features.</p>', unsafe_allow_html=True)
    st.markdown('<div class="section-rule"></div>', unsafe_allow_html=True)
    st.markdown(
        '<p class="sidebar-note">Sources and indexes stay in <code>data/</code>. '
        'Do not commit private research papers to GitHub.</p>',
        unsafe_allow_html=True,
    )


def section_header(kicker: str, title: str, copy: str) -> None:
    st.markdown(f'<div class="eyebrow">{kicker}</div>', unsafe_allow_html=True)
    st.title(title)
    st.markdown(f'<p class="hero-copy">{copy}</p>', unsafe_allow_html=True)


def show_document_rows(items: list[dict], limit: int = 20) -> None:
    if not items:
        st.info("No indexed sources yet. Add PDF, Word, Excel, PowerPoint, image, TXT, or CSV files from the library.")
        return
    for item in items[:limit]:
        with st.container(border=True):
            left, middle, right = st.columns([5, 1.3, 1.2])
            left.markdown(f"**{item['filename']}**")
            left.caption(f"{item['file_type'].upper()} · {item['chunk_count']} source passages")
            middle.markdown(f"`{item['file_type'].upper()}`")
            if right.button("Remove", key=f"delete-{item['document_id']}"):
                library.delete_document(item["document_id"])
                st.rerun()
    if len(items) > limit:
        st.caption(f"Showing {limit} of {len(items)} sources. Use search to narrow the collection.")


if page == "Home":
    section_header(
        "SECOND BRAIN / FIELDNOTES",
        "Start with a question. Write with evidence.",
        "A local research workspace to explore your literature, connect findings across papers, and build a source-grounded draft.",
    )
    with st.container(horizontal=True):
        st.metric("Papers in library", f"{len(documents):,}", border=True)
        st.metric("Indexed passages", f"{passage_count:,}", border=True)
        known_countries = sum(bool(document.get("country")) for document in documents)
        st.metric("Affiliation countries detected", f"{known_countries:,}", border=True)
    left, right = st.columns([1.45, 1])
    with left:
        with st.container(border=True):
            st.markdown('<div class="eyebrow">RESEARCH ASSISTANT</div>', unsafe_allow_html=True)
            st.subheader("What are you exploring?")
            with st.form("home_research_form"):
                home_question = st.text_input(
                    "Ask your research library",
                    placeholder="e.g. How are dynamic tests used to update bridge models?",
                )
                submitted = st.form_submit_button(
                    "Ask Qwen",
                    type="primary",
                    icon=":material/arrow_forward:",
                )
            if submitted and home_question.strip():
                chunks, retrieval_mode = library.search(
                    home_question, ollama_url, embedding_model, limit=7
                )
                if chunks:
                    try:
                        st.session_state.home_answer = ask_ollama(
                            ollama_url,
                            text_model,
                            keep_alive,
                            "research library",
                            format_context(chunks),
                            [],
                            home_question,
                            "English",
                        )
                        st.session_state.home_sources = list(
                            dict.fromkeys(
                                f"{chunk['source']} · {chunk['locator']}" for chunk in chunks
                            )
                        )
                        st.session_state.home_retrieval_mode = retrieval_mode
                    except (requests.RequestException, KeyError, ValueError) as error:
                        st.error(f"Qwen could not answer: {error}")
                else:
                    st.warning("No matching source passages. Try a different search.")
            if st.session_state.get("home_answer"):
                st.markdown(st.session_state.home_answer)
                st.caption(
                    f"{st.session_state.get('home_retrieval_mode', 'keyword').title()} search · "
                    + " · ".join(st.session_state.get("home_sources", []))
                )
            st.button(
                "Explore library",
                on_click=select_workspace,
                args=("Library",),
                icon=":material/library_books:",
            )
    with right:
        st.markdown('<div class="eyebrow">FIELD NOTES</div>', unsafe_allow_html=True)
        st.subheader("Recently added")
        if documents:
            for item in documents[:5]:
                title = item.get("title") or item["filename"]
                st.markdown(f"**{title[:90]}**")
                location = item.get("country") or "Location unverified"
                year = item.get("year") or "Year unknown"
                st.caption(f"{year} · {location} · {item['chunk_count']} passages")
        else:
            st.info("Your local library is empty.")
        st.button(
            "Open page builder",
            on_click=select_workspace,
            args=("Page builder",),
            icon=":material/edit_document:",
        )

elif page == "Library":
    if not st.session_state.get("bibliography_backfilled"):
        with st.spinner("Reading first-page bibliography metadata locally..."):
            metadata_updated, metadata_failed = library.backfill_metadata()
        st.session_state.bibliography_backfilled = True
        if metadata_updated:
            st.toast(f"Checked metadata for {metadata_updated} papers.")
        if metadata_failed:
            st.caption(f"Could not read bibliography metadata from {metadata_failed} files.")
        documents = library.documents()
    section_header(
        "01 · LITERATURE LIBRARY",
        "Explore your research corpus.",
        "Search paper text and keywords, filter by publication year or affiliation country, and review detected metadata before exploring bibliometric patterns.",
    )
    reference_dir = ROOT / "Referensi"
    reference_pdfs = sorted(path for path in reference_dir.glob("*") if path.suffix.casefold() == ".pdf") if reference_dir.exists() else []
    if reference_pdfs:
        with st.expander(f"Import local reference folder ({len(reference_pdfs)} PDFs)"):
            with st.form("folder_import_form"):
                st.caption("PDF text is imported locally by default. Scanned PDFs require Qwen-VL and may take longer.")
                analyze_scans = st.checkbox(
                    "Analyze up to two scanned pages with Qwen-VL",
                    value=False,
                )
                import_folder = st.form_submit_button(
                    "Index local PDFs",
                    icon=":material/folder_open:",
                )
            if import_folder:
                added_count = 0
                skipped_count = 0
                failures = []
                progress = st.progress(0)
                for index, path in enumerate(reference_pdfs, 1):
                    try:
                        raw_data = path.read_bytes()
                        if library.contains_document(raw_data):
                            skipped_count += 1
                            progress.progress(index / len(reference_pdfs))
                            continue
                        extracted = extract_document(
                            path.name,
                            raw_data,
                            ollama_url,
                            vision_model if analyze_scans else "",
                        )
                        _, added, _ = library.add_document(extracted, raw_data, ollama_url)
                        added_count += int(added)
                    except Exception as error:
                        failures.append(f"{path.name}: {error}")
                    progress.progress(index / len(reference_pdfs))
                progress.empty()
                st.success(f"Indexed {added_count} new PDFs; skipped {skipped_count} already in the library.")
                for failure in failures[:8]:
                    st.warning(failure)
                if len(failures) > 8:
                    st.caption(f"{len(failures) - 8} additional errors not shown.")
                documents = library.documents()
                all_chunks = library.all_chunks()

    with st.expander("Add other sources", icon=":material/upload_file:"):
        uploaded_files = st.file_uploader(
            "Choose one or more files",
            type=[extension.lstrip(".") for extension in sorted(SUPPORTED_EXTENSIONS)],
            accept_multiple_files=True,
            help="PDF, DOCX, XLSX, PPTX, images, TXT, and CSV.",
        )
        inspect_visuals = st.checkbox("Analyze charts/diagrams with Qwen-VL", value=False)
        if st.button("Index selected files", type="primary", disabled=not uploaded_files):
            added_count = 0
            failures = []
            progress = st.progress(0)
            for index, uploaded_file in enumerate(uploaded_files, 1):
                try:
                    raw_data = uploaded_file.getvalue()
                    extracted = extract_document(
                        uploaded_file.name,
                        raw_data,
                        ollama_url,
                        vision_model,
                        inspect_visuals,
                    )
                    _, added, note = library.add_document(
                        extracted, raw_data, ollama_url, embedding_model
                    )
                    added_count += int(added)
                    if "unavailable" in note:
                        failures.append(f"{uploaded_file.name}: {note}")
                except Exception as error:
                    failures.append(f"{uploaded_file.name}: {error}")
                progress.progress(index / len(uploaded_files))
            progress.empty()
            st.success(f"Indexed {added_count} new sources.")
            for failure in failures[:8]:
                st.warning(failure)
            documents = library.documents()
            all_chunks = library.all_chunks()

    country_options = sorted(COUNTRY_LOCATIONS)
    years = sorted({int(document["year"]) for document in documents if document.get("year")})
    library_query = st.text_input(
        "Search paper text, titles, or keywords",
        placeholder="e.g. finite element model updating",
        key="library_fulltext_query",
    )
    country_filter_col, year_filter_col = st.columns(2)
    with country_filter_col:
        selected_countries = st.multiselect(
            "Affiliation country",
            ["Unknown", *country_options],
            placeholder="All countries",
        )
    with year_filter_col:
        year_range = None
        if len(years) > 1:
            year_range = st.slider("Publication year", min(years), max(years), (min(years), max(years)))
        elif years:
            year_range = (years[0], years[0])
            st.caption(f"Publication year: {years[0]}")

    visible_documents = filter_documents(
        documents,
        all_chunks,
        query=library_query,
        countries=["Belum diketahui" if country == "Unknown" else country for country in selected_countries],
        year_range=year_range,
    )
    visible_ids = {document["document_id"] for document in visible_documents}
    visible_names = {document["filename"] for document in visible_documents}
    visible_chunks = [chunk for chunk in all_chunks if chunk["source"] in visible_names]
    metrics = bibliometric_summary(visible_documents, visible_chunks)

    with st.container(horizontal=True):
        st.metric("Matching papers", f"{len(visible_documents):,}", border=True)
        st.metric("Publication span", f"{min(years)}–{max(years)}" if years else "Unknown", border=True)
        st.metric("Affiliation countries", f"{len(metrics['country_counts'])}", border=True)
        st.metric("Countries to review", f"{metrics['unknown_country_count']:,}", border=True)

    trend_col, map_col = st.columns([1, 1.15])
    with trend_col:
        st.subheader("Publication trend")
        if metrics["year_counts"]:
            trend_frame = pd.DataFrame(metrics["year_counts"], columns=["Year", "Papers"])
            trend_color = "#2E766F" if active_theme == "light" else "#79C0B5"
            trend = (
                alt.Chart(trend_frame)
                .mark_area(color=trend_color, opacity=0.2)
                .encode(
                    x=alt.X("Year:Q", title="Year", axis=alt.Axis(format="d", tickCount=7)),
                    y=alt.Y("Papers:Q", title="Papers", scale=alt.Scale(zero=True)),
                    tooltip=["Year:Q", "Papers:Q"],
                )
            )
            line = trend.mark_line(color=trend_color, point=alt.OverlayMarkDef(filled=True, size=34))
            st.altair_chart((trend + line).properties(height=245), width="stretch")
        else:
            st.caption("No publication years detected in the matching sources.")
    with map_col:
        st.subheader("Paper affiliation map")
        if metrics["map_points"]:
            map_data = pd.DataFrame(metrics["map_points"])
            map_data["marker_size"] = map_data["publications"].map(
                lambda count: max(18000, min(80000, count * 9000))
            )
            st.map(
                map_data,
                latitude="latitude",
                longitude="longitude",
                size="marker_size",
                color="#BB5A3C" if active_theme == "light" else "#F08A69",
                height=360,
            )
            st.caption(
                "Pins show detected or manually verified country centroids, not experiment sites. "
                "Map tiles are served by Carto; paper text is never sent to the map provider."
            )
        else:
            st.caption("No affiliation countries detected for this selection.")

    keyword_col, country_col = st.columns(2)
    with keyword_col:
        st.subheader("Frequent research themes")
        st.caption(f"Theme source: {metrics['keyword_source']}.")
        if metrics["keyword_counts"]:
            keyword_frame = pd.DataFrame(metrics["keyword_counts"], columns=["Term", "Papers"])
            keyword_chart = (
                alt.Chart(keyword_frame)
                .mark_bar(color="#CC6B4E" if active_theme == "light" else "#F08A69", cornerRadiusEnd=3)
                .encode(
                    x=alt.X("Papers:Q", title="Papers"),
                    y=alt.Y("Term:N", sort="-x", title=None, axis=alt.Axis(labelLimit=160)),
                    tooltip=["Term:N", "Papers:Q"],
                )
                .properties(height=270)
            )
            st.altair_chart(keyword_chart, width="stretch")
        else:
            st.caption("No bibliographic keywords available.")
    with country_col:
        st.subheader("Papers by affiliation country")
        if metrics["country_counts"]:
            country_frame = pd.DataFrame(metrics["country_counts"], columns=["Country", "Papers"])
            country_chart = (
                alt.Chart(country_frame)
                .mark_bar(color="#547F9F" if active_theme == "light" else "#82AFC8", cornerRadiusEnd=3)
                .encode(
                    x=alt.X("Papers:Q", title="Papers"),
                    y=alt.Y("Country:N", sort="-x", title=None, axis=alt.Axis(labelLimit=150)),
                    tooltip=["Country:N", "Papers:Q"],
                )
                .properties(height=270)
            )
            st.altair_chart(country_chart, width="stretch")
        else:
            st.caption("No affiliation countries detected.")

    st.markdown("### Bibliographic metadata · editable")
    st.caption("Detected countries come from first-page affiliation text. Verify them before collaboration analysis.")
    if visible_documents:
        editor_rows = [
            {
                "document_id": document["document_id"],
                "File": document["filename"],
                "Title": document.get("title") or "",
                "Year": document.get("year"),
                "Affiliation country": document.get("country") or "",
                "Country status": (
                    "Manually verified"
                    if document.get("country_source") == "manually verified"
                    else "Detected · review"
                    if document.get("country")
                    else "Unknown"
                ),
                "Author keywords · separate with ;": document.get("keywords") or "",
            }
            for document in visible_documents
        ]
        editor_frame = pd.DataFrame(editor_rows)
        editor_frame["Year"] = pd.array(editor_frame["Year"], dtype="Int64")
        editor_key = "bibliography_editor_" + str(hash(tuple(sorted(visible_ids))))
        edited_frame = st.data_editor(
            editor_frame,
            key=editor_key,
            hide_index=True,
            width="stretch",
            height=420,
            num_rows="fixed",
            column_config={
                "document_id": None,
                "File": st.column_config.TextColumn("File", disabled=True),
                "Title": st.column_config.TextColumn("Paper title"),
                "Year": st.column_config.NumberColumn("Year", min_value=1800, max_value=2100, step=1),
                "Affiliation country": st.column_config.SelectboxColumn(
                    "Affiliation country",
                    options=["", *country_options],
                ),
                "Country status": st.column_config.TextColumn("Country status", disabled=True),
                "Author keywords · separate with ;": st.column_config.TextColumn("Author keywords · separate with ;"),
            },
        )
        if st.button("Save metadata corrections", icon=":material/save:"):
            original_by_id = {document["document_id"]: document for document in visible_documents}
            saved_count = 0
            for _, row in edited_frame.iterrows():
                document_id = row["document_id"]
                original = original_by_id[document_id]
                year_value = row["Year"]
                year_value = None if pd.isna(year_value) else int(year_value)
                title_value = str(row["Title"] or "")
                country_value = str(row["Affiliation country"] or "")
                keywords_value = str(row["Author keywords · separate with ;"] or "")
                if (
                    title_value != (original.get("title") or "")
                    or year_value != original.get("year")
                    or country_value != (original.get("country") or "")
                    or keywords_value != (original.get("keywords") or "")
                ):
                    library.update_metadata(
                        document_id,
                        title=title_value,
                        year=year_value,
                        country=country_value,
                        keywords=keywords_value,
                    )
                    saved_count += 1
            st.toast(f"Saved {saved_count} metadata corrections.")
            st.rerun()
    else:
        st.info("No papers match these filters.")

elif page == "Ask library":
    section_header(
        "02 · EVIDENCE AND SYNTHESIS",
        "Ask your research library.",
        "Answers use passages retrieved from your local corpus. Check the file and page labels before using a claim in your writing.",
    )
    if st.button("Clear conversation"):
        st.session_state.research_messages = []
        st.rerun()
    messages = st.session_state.setdefault("research_messages", [])
    for message in messages:
        with st.chat_message(message["role"]):
            st.markdown(message["content"])
            if message.get("sources"):
                st.caption("Sources: " + " · ".join(message["sources"]))
    question = st.chat_input(
        "Ask a question about your sources",
        disabled=not documents,
    )
    if not documents:
        st.info("Add sources in the Library before asking a question.")
    if question:
        messages.append({"role": "user", "content": question})
        with st.chat_message("user"):
            st.markdown(question)
        with st.chat_message("assistant"):
            with st.spinner("Searching local sources and preparing an answer..."):
                chunks, retrieval_mode = library.search(
                    question, ollama_url, embedding_model, limit=7
                )
                if not chunks:
                    answer = "I couldn't find matching source passages. Try different terms or check that the papers are indexed."
                    citations = []
                else:
                    context = format_context(chunks)
                    try:
                        answer = ask_ollama(
                            ollama_url,
                            text_model,
                            keep_alive,
                            "research library",
                            context,
                            messages[:-1],
                            question,
                            "English",
                        )
                    except (requests.RequestException, KeyError, ValueError) as error:
                        answer = f"Could not get an answer from Ollama: {error}"
                    citations = list(
                        dict.fromkeys(f"{chunk['source']} · {chunk['locator']}" for chunk in chunks)
                    )
            st.markdown(answer)
            if citations:
                st.caption(f"{retrieval_mode.title()} search · Sources: " + " · ".join(citations))
        messages.append({"role": "assistant", "content": answer, "sources": citations})

elif page == "Page builder":
    section_header(
        "03 · PAGE BUILDER",
        "Write with your sources in view.",
        "Build a paper one section at a time, preview the page beside your text, and ask Qwen to refine wording without dropping citations.",
    )
    saved_draft = library.load_draft()
    for old_name, new_name in LEGACY_SECTION_NAMES.items():
        if old_name in saved_draft["sections"]:
            saved_draft["sections"].setdefault(new_name, saved_draft["sections"].pop(old_name))
        if old_name in saved_draft["sources"]:
            saved_draft["sources"].setdefault(new_name, saved_draft["sources"].pop(old_name))
    draft_sections = st.session_state.setdefault("draft_sections", saved_draft["sections"])
    draft_sources = st.session_state.setdefault("draft_sources", saved_draft["sources"])
    st.session_state.setdefault("draft_title", saved_draft["title"])
    st.session_state.setdefault("research_question", saved_draft["research_question"])
    st.session_state.setdefault("research_notes", saved_draft["research_notes"])
    title = st.text_input(
        "Working title",
        key="draft_title",
        placeholder="Enter a working title",
        on_change=save_current_draft,
    )
    research_question = st.text_input(
        "Research question or objective",
        key="research_question",
        placeholder="What should this study answer?",
        on_change=save_current_draft,
    )
    section_name = st.selectbox("Section to write", SECTION_ORDER, index=1)
    research_notes = st.text_area(
        "Researcher notes",
        height=110,
        key="research_notes",
        on_change=save_current_draft,
        placeholder="Add your actual design, data, methods, or findings. Results are never invented from literature alone.",
    )
    if st.button("Draft this section", type="primary", disabled=not documents):
        search_query = f"{section_name}. {title}. {research_question}. {research_notes}"
        chunks, retrieval_mode = library.search(
            search_query, ollama_url, embedding_model, limit=8
        )
        if not chunks:
            st.warning("No matching sources for this topic yet.")
        else:
            prompt = (
                f"Write only a first draft in English for the research paper section '{section_name}'. "
                f"Working title: {title or 'not provided'}. Research question: "
                f"{research_question or 'not provided'}. Researcher notes and actual data: "
                f"{research_notes or 'none provided'}. Synthesize rather than closely "
                "paraphrase any one source. Cite claims with exact source labels copied "
                "from the excerpts. Never invent findings, procedures, sample sizes, "
                "measurements, or references. For missing study-specific information, "
                "insert [Researcher input needed]. Return prose only."
            )
            try:
                draft = ask_ollama(
                    ollama_url,
                    text_model,
                    keep_alive,
                    "research library",
                    format_context(chunks),
                    [],
                    prompt,
                    "English",
                )
                draft_sections[section_name] = draft
                draft_sources[section_name] = list(
                    dict.fromkeys(
                        f"{chunk['source']} · {chunk['locator']}" for chunk in chunks
                    )
                )
                st.session_state[f"pending-draft-{section_name}"] = draft
                st.session_state.draft_search_mode = retrieval_mode
                library.save_draft(
                    title=title,
                    research_question=research_question,
                    research_notes=research_notes,
                    sections=draft_sections,
                    sources=draft_sources,
                )
                st.rerun()
            except (requests.RequestException, KeyError, ValueError) as error:
                st.error(f"Could not create the draft: {error}")

    st.markdown('<div class="section-rule"></div>', unsafe_allow_html=True)
    if not documents:
        st.info("Add and index sources before drafting from the literature.")
    editor_column, preview_column = st.columns([1, 1])
    editor_key = f"draft-edit-buffer-{section_name}"
    pending_draft_key = f"pending-draft-{section_name}"
    if pending_draft_key in st.session_state:
        st.session_state[editor_key] = st.session_state.pop(pending_draft_key)
    else:
        st.session_state.setdefault(editor_key, draft_sections.get(section_name, ""))
    with editor_column:
        st.subheader("Manuscript canvas")
        st.caption("Apply your edits to refresh the page preview and enable Qwen actions.")
        with st.form(f"draft-editor-form-{section_name}", border=False):
            edited_text = st.text_area(
                f"{section_name} draft text",
                height=420,
                key=editor_key,
                label_visibility="collapsed",
                placeholder="Your section draft will appear here, or start writing manually.",
            )
            apply_draft = st.form_submit_button(
                "Apply changes",
                icon=":material/check:",
            )
        if apply_draft:
            draft_sections[section_name] = edited_text
            save_current_draft()
            st.toast("Draft saved to the local database.")
            st.rerun()

        context_labels = set(draft_sources.get(section_name, []))
        source_chunks = [
            chunk
            for chunk in all_chunks
            if f"{chunk['source']} · {chunk['locator']}" in context_labels
        ][:8]
        source_context = format_context(source_chunks)
        action_left, action_right = st.columns([1, 1])
        with action_left:
            if st.button(
                "Paraphrase with Qwen",
                key=f"paraphrase-{section_name}",
                icon=":material/auto_awesome:",
                disabled=not draft_sections.get(section_name, "").strip(),
            ):
                current_text = draft_sections.get(section_name, "")
                try:
                    rewritten = ask_ollama(
                        ollama_url,
                        text_model,
                        keep_alive,
                        "research manuscript",
                        source_context,
                        [],
                        "Paraphrase and polish the following text in natural academic English. "
                        "Preserve the meaning and every bracketed citation exactly; do not add "
                        "claims, numbers, or references. If the source evidence is insufficient, "
                        "keep the marker [Researcher input needed]. Return only the revised text.\n\n"
                        f"Text:\n{current_text}",
                        "English",
                    )
                    draft_sections[section_name] = rewritten
                    st.session_state[pending_draft_key] = rewritten
                    save_current_draft()
                    st.rerun()
                except (requests.RequestException, KeyError, ValueError) as error:
                    st.error(f"Paraphrasing failed: {error}")
        with action_right:
            qwen_question = st.text_input(
                "Ask Qwen about this section",
                key=f"builder-question-{section_name}",
                placeholder="e.g. Which claims need stronger sources?",
            )
            if st.button(
                "Ask about this text",
                key=f"ask-builder-{section_name}",
                icon=":material/forum:",
                disabled=not qwen_question.strip(),
            ):
                current_text = draft_sections.get(section_name, "")
                try:
                    answer = ask_ollama(
                        ollama_url,
                        text_model,
                        keep_alive,
                        "research manuscript",
                        source_context,
                        [],
                        f"Answer this question about the draft using source context where available. "
                        f"Include complete source labels: {qwen_question}\n\nDraft:\n{current_text}",
                        "English",
                    )
                    st.session_state[f"builder-answer-{section_name}"] = answer
                except (requests.RequestException, KeyError, ValueError) as error:
                    st.error(f"Qwen could not answer: {error}")
        if st.session_state.get(f"builder-answer-{section_name}"):
            st.info(st.session_state[f"builder-answer-{section_name}"], icon=":material/auto_awesome:")
        st.caption("Drafts autosave locally when you apply changes.")
        citations = draft_sources.get(section_name, [])
        if citations:
            st.caption("Sources used: " + " · ".join(citations))
        st.caption("Autosaved to the local database.")
    with preview_column:
        st.subheader("Page preview")
        escaped_title = html.escape(title or "Research paper title")
        escaped_section = html.escape(section_name)
        escaped_text = html.escape(
            draft_sections.get(section_name, "") or "Your section text will appear here."
        )
        st.markdown(
            f'<div class="paper-preview"><div class="paper-meta">RESEARCH DRAFT · {escaped_section.upper()}</div>'
            f'<h3>{escaped_title}</h3><p>{escaped_text}</p></div>',
            unsafe_allow_html=True,
        )

    existing_sections = {key: value for key, value in draft_sections.items() if value.strip()}
    if existing_sections:
        st.markdown('<div class="section-rule"></div>', unsafe_allow_html=True)
        docx_data = make_docx(title, existing_sections, draft_sources)
        st.download_button(
            "Export manuscript to Word",
            data=docx_data,
            file_name="second-brain-research-draft.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            type="primary",
        )
        st.caption("DOCX is generated locally. Source labels still need final citation-style formatting.")